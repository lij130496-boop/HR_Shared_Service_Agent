from __future__ import annotations

from io import BytesIO
from datetime import datetime

from docx import Document

from docx.shared import (
    Pt,
    Cm
)

from docx.enum.text import (
    WD_ALIGN_PARAGRAPH
)

from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT
)

from docx.oxml import OxmlElement
from docx.oxml.ns import qn



# ======================================================
# 创建基础文档
# ======================================================

def _new_doc(title: str):

    doc = Document()


    section = doc.sections[0]

    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)



    style = doc.styles["Normal"]

    style.font.name = "宋体"

    style.font.size = Pt(11)

    style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "宋体"
    )


    p = doc.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    r = p.add_run(
        title
    )

    r.bold = True

    r.font.size = Pt(20)

    rPr = r._element.get_or_add_rPr()

    rFonts = rPr.get_or_add_rFonts()

    rFonts.set(
        qn("w:eastAsia"),
        "黑体"
    )


    return doc



# ======================================================
# 保存
# ======================================================

def _finish(doc):

    buf = BytesIO()

    doc.save(buf)

    return buf.getvalue()



# ======================================================
# 表格工具
# ======================================================

def set_cell(
    cell,
    text,
    bold=False
):

    cell.text = ""

    p = cell.paragraphs[0]

    r = p.add_run(
        str(text)
    )

    r.bold = bold

    r.font.size = Pt(10)

    rPr = r._element.get_or_add_rPr()

    rFonts = rPr.get_or_add_rFonts()

    rFonts.set(
        qn("w:eastAsia"),
        "宋体"
    )

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )



def gray(cell):

    tcPr = (
        cell._tc
        .get_or_add_tcPr()
    )

    shd = OxmlElement(
        "w:shd"
    )

    shd.set(
        qn("w:fill"),
        "EAEAEA"
    )

    tcPr.append(shd)



# ======================================================
# 主函数
# ======================================================


def generate_document(
    document_type: str,
    data: dict
):


    # ==================================================
    # 1. 请假申请
    # ==================================================

    if document_type == "leave_request":


        doc = _new_doc(
            "员工请假申请单"
        )


        p = doc.add_paragraph()

        p.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT
        )

        p.add_run(
            "填表日期："
            +
            datetime.now()
            .strftime("%Y-%m-%d")
        )



        table = doc.add_table(
            rows=5,
            cols=6
        )

        table.style = (
            "Table Grid"
        )


        values = [

            [
                "请假人",
                data.get("name",""),
                "部门",
                data.get("department",""),
                "职务",
                data.get("position","")
            ],


            [
                "请假类别",

                "□事假  □病假  □婚假  □丧假  □产假",

                "",
                "",
                "",
                ""
            ],


            [
                "请假事由",

                data.get(
                    "reason",
                    ""
                ),

                "",
                "",
                "",
                ""
            ],


            [
                "请假时间",

                f"{data.get('start_date','')} 至 {data.get('end_date','')}",

                "",
                "",
                "",
                ""
            ],


            [
                "审批意见",

                "部门主管意见：\n\n签字：________",

                "总经理意见：\n\n签字：________",

                "",

                "人事审批：\n\n签字：________",

                ""
            ]

        ]


        for i,row in enumerate(values):

            for j,value in enumerate(row):

                set_cell(
                    table.cell(i,j),
                    value,
                    j in [0,2,4]
                )

                if j in [0,2,4]:

                    gray(
                        table.cell(i,j)
                    )


        filename = (
            f"{data.get('name','员工')}"
            "_请假申请单.docx"
        )



    # ==================================================
    # 2. 加班申请
    # ==================================================

    elif document_type == "overtime_request":


        doc = _new_doc(
            "员工加班申请单"
        )


        table = doc.add_table(
            rows=5,
            cols=2
        )

        table.style="Table Grid"


        rows=[

            ("申请人",
             data.get("name","")),

            ("部门",
             data.get("department","")),

            ("加班日期",
             data.get("work_date","")),

            ("预计时长",
             data.get("hours","")),

            ("加班原因",
             data.get("reason",""))

        ]


        for i,(a,b) in enumerate(rows):

            set_cell(
                table.cell(i,0),
                a,
                True
            )

            gray(
                table.cell(i,0)
            )

            set_cell(
                table.cell(i,1),
                b
            )


        filename=f"{data.get('name','员工')}_加班申请.docx"



    # ==================================================
    # 3. 面试邀请
    # ==================================================

    elif document_type=="interview_invitation":


        name=data.get(
            "candidate",
            "候选人"
        )


        doc=_new_doc(
            "面试邀请函"
        )


        doc.add_paragraph(
            f"{name}先生/女士，您好："
        )


        doc.add_paragraph(

            f"感谢您关注我公司"
            f"{data.get('position','')}"
            f"岗位，诚邀您参加面试。"
        )


        table=doc.add_table(
            rows=4,
            cols=2
        )

        table.style="Table Grid"


        info=[

            ("面试时间",
             data.get("interview_time","")),

            ("面试地点",
             data.get("location","")),

            ("联系人",
             data.get("contact","")),

            ("联系电话",
             data.get("phone",""))

        ]


        for i,(a,b) in enumerate(info):

            set_cell(
                table.cell(i,0),
                a,
                True
            )

            gray(
                table.cell(i,0)
            )

            set_cell(
                table.cell(i,1),
                b
            )


        filename=f"{name}_面试邀请函.docx"



    # ==================================================
    # 4. 入职通知
    # ==================================================

    elif document_type=="onboarding_notice":


        doc=_new_doc(
            "员工入职通知书"
        )


        name=data.get(
            "name",
            "新员工"
        )


        doc.add_paragraph(
            f"{name}您好："
        )


        doc.add_paragraph(

            "欢迎加入公司。"
            "请按照以下时间办理入职手续。"
        )


        for k,v in [

            ("入职部门",
             data.get("department")),

            ("岗位",
             data.get("position")),

            ("报到时间",
             data.get("report_time")),

            ("报到地点",
             data.get("location"))

        ]:

            doc.add_paragraph(
                f"{k}：{v}"
            )


        filename=f"{name}_入职通知.docx"



    # ==================================================
    # 5. 在职证明
    # ==================================================

    elif document_type=="employment_certificate":


        doc=_new_doc(
            "在职证明"
        )


        name=data.get(
            "name",
            ""
        )


        doc.add_paragraph(

            f"兹证明{name}为我公司正式员工，"
            f"现任{data.get('department','')} "
            f"{data.get('position','')}岗位。"
        )


        doc.add_paragraph(
            "特此证明。"
        )


        doc.add_paragraph(
            "日期："
            +
            datetime.now()
            .strftime("%Y-%m-%d")
        )


        filename=f"{name}_在职证明.docx"



    # ==================================================
    # 6. 培训通知
    # ==================================================

    elif document_type=="training_notice":


        doc=_new_doc(
            "培训通知"
        )


        for k,v in [

            ("培训主题",
             data.get("topic")),

            ("培训时间",
             data.get("training_time")),

            ("培训地点",
             data.get("location")),

            ("培训对象",
             data.get("audience")),

            ("培训讲师",
             data.get("lecturer"))

        ]:

            doc.add_paragraph(
                f"{k}：{v}"
            )


        filename=f"{data.get('topic','培训')}_培训通知.docx"



    # ==================================================
    # 7. 会议纪要
    # ==================================================

    elif document_type=="meeting_minutes":


        doc=_new_doc(
            "会议纪要"
        )


        for k,v in [

            ("会议主题",
             data.get("topic")),

            ("会议时间",
             data.get("meeting_time")),

            ("参会人员",
             data.get("participants"))

        ]:

            doc.add_paragraph(
                f"{k}：{v}"
            )


        doc.add_heading(
            "会议内容",
            level=2
        )

        doc.add_paragraph(
            data.get(
                "discussion",
                ""
            )
        )


        doc.add_heading(
            "会议决定",
            level=2
        )

        doc.add_paragraph(
            data.get(
                "decisions",
                ""
            )
        )


        filename=f"{data.get('topic','会议')}_会议纪要.docx"



    else:

        raise ValueError(
            f"不支持的文档类型:{document_type}"
        )



    return (
        _finish(doc),
        filename
    )