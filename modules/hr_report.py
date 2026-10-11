from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document


def _read_jsonl(path: Path):
    if not path.exists():
        return []
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            rows.append(json.loads(line))
        except Exception:
            pass
    return rows


def build_hr_monthly_report(base_dir: Path) -> tuple[bytes, str]:
    data_dir = base_dir / "data"

    now = datetime.now()
    month_prefix = f"{now.year}-{now.month:02d}"

    def in_current_month(item: dict) -> bool:
        time_text = str(item.get("time") or item.get("submitted_at") or "")
        return time_text.startswith(month_prefix)

    questions = [
        x for x in _read_jsonl(data_dir / "ai_questions.jsonl")
        if in_current_month(x)
    ]
    feedbacks = [
        x for x in _read_jsonl(data_dir / "feedback.jsonl")
        if in_current_month(x)
    ]
    hitl = [
        x for x in _read_jsonl(data_dir / "hitl_log.jsonl")
        if in_current_month(x)
    ]
    onboarding = [
        x for x in _read_jsonl(data_dir / "onboarding_records.jsonl")
        if in_current_month(x)
    ]

    good = sum(1 for x in feedbacks if x.get("status") == "success")
    bad = sum(1 for x in feedbacks if x.get("status") == "failed")
    total_feedback = good + bad
    satisfaction = (good / total_feedback * 100) if total_feedback else 0

    failed_questions = [x for x in questions if x.get("status") == "failed"]

    counter = Counter(
        (x.get("question") or "").strip()
        for x in questions
        if (x.get("question") or "").strip()
    )

    title = f"{now.year}年{now.month:02d}月 HR智能运营报告"

    doc = Document()
    doc.add_heading(title, level=1)

    doc.add_heading("一、核心数据", level=2)
    doc.add_paragraph(f"员工咨询总量：{len(questions)}")
    doc.add_paragraph(f"员工反馈总量：{total_feedback}")
    doc.add_paragraph(f"用户认可率：{satisfaction:.1f}%")
    doc.add_paragraph(f"待人工确认记录：{len(hitl)}")
    doc.add_paragraph(f"新员工培训完成：{len(onboarding)}")

    doc.add_heading("二、高频问题", level=2)
    if counter:
        for q, n in counter.most_common(10):
            doc.add_paragraph(f"{q}（{n}次）")
    else:
        doc.add_paragraph("暂无咨询数据。")

    doc.add_heading("三、待优化问题", level=2)
    if failed_questions:
        for item in failed_questions[:20]:
            doc.add_paragraph(item.get("question", "-"))
    else:
        doc.add_paragraph("本期暂无明显未解决问题。")

    doc.add_heading("四、建议", level=2)
    doc.add_paragraph("建议HR重点补充高频未解决问题对应的制度资料，并持续关注员工负面反馈和人工审核记录。")

    buf = BytesIO()
    doc.save(buf)

    return buf.getvalue(), f"{now.year}-{now.month:02d}_HR智能运营报告.docx"
